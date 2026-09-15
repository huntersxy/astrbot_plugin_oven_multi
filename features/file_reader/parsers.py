# Modified from astrbot_plugin_file_reader_pro (MIT) by zz6zz666
# https://github.com/zz6zz666/astrbot_plugin_file_reader_pro
"""文件内容解析：将各类文档解析为纯文本。

相对原版的调整：
- 错误以异常抛出，不再与正文内容混在同一个返回值里；
- .doc/.ppt 等旧二进制格式明确拒绝（python-docx/python-pptx 无法解析，原版的转换逻辑实际不可用）；
- MIME 检测（python-magic）为可选增强，未安装时自动退回扩展名判断；
- CSV 先用 chardet 检测编码，避免中文文件乱码。
"""

import os
from typing import Callable, Dict, Optional

import chardet


class FileParseError(Exception):
    """文件解析失败，message 为可直接展示给用户的说明。"""


# 按纯文本读取的扩展名
TEXT_EXTENSIONS = {
    # 文档/标记语言
    "md", "markdown", "html", "htm", "xml", "rst", "adoc",
    # 配置与数据描述
    "json", "yaml", "yml", "ini", "cfg", "conf", "properties", "env", "toml", "lock",
    # 编程语言
    "py", "java", "cpp", "c", "h", "hpp", "cs", "js", "ts", "jsx", "tsx", "php",
    "rb", "go", "rs", "swift", "kt", "scala", "sh", "bash", "ps1", "bat", "cmd", "vbs", "sql",
    # 其他文本
    "txt", "log", "csv", "tsv", "url", "gitignore",
    "",  # 无扩展名文件按文本处理
}

# 需要专用解析器的扩展名
SPECIAL_EXTENSIONS = {"pdf", "docx", "xlsx", "xls", "ods", "pptx", "odp"}

# MIME → 扩展名的确定映射（python-magic 可用时优先于扩展名）
_MIME_TO_EXT = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/plain": "txt",
    "text/csv": "csv",
    "text/markdown": "md",
    "application/json": "json",
    "application/xml": "xml",
    "text/xml": "xml",
    "text/html": "html",
    "application/zip": "zip",
}


def detect_extension(file_path: str) -> str:
    """检测文件类型（优先 MIME，退回扩展名），返回小写、不带点的扩展名。"""
    if not os.path.isfile(file_path):
        raise FileParseError(f"文件不存在: {file_path}")

    try:
        import magic

        mime = magic.from_file(file_path, mime=True) or ""
        if mime in _MIME_TO_EXT:
            return _MIME_TO_EXT[mime]
    except ImportError:
        pass
    except Exception:
        pass  # MIME 检测失败时退回扩展名

    ext = os.path.splitext(file_path)[1]
    return ext[1:].lower() if ext else ""


def _read_text(file_path: str) -> str:
    """读取文本文件，自动检测编码。"""
    with open(file_path, "rb") as f:
        raw = f.read()
    encoding = (chardet.detect(raw) or {}).get("encoding") or "utf-8"
    try:
        text = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        text = raw.decode("utf-8", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _read_csv(file_path: str) -> str:
    """读取 CSV/TSV 文件。"""
    import pandas as pd

    with open(file_path, "rb") as f:
        raw = f.read(65536)
    encoding = (chardet.detect(raw) or {}).get("encoding") or "utf-8"
    sep = "\t" if file_path.lower().endswith(".tsv") else ","
    try:
        df = pd.read_csv(file_path, encoding=encoding, sep=sep)
    except Exception:
        df = pd.read_csv(file_path, sep=sep)  # 编码误判时让 pandas 自行处理
    return df.to_string(index=False)


def _read_pdf(file_path: str) -> str:
    from pdfminer.high_level import extract_text

    return extract_text(file_path)


def _read_docx(file_path: str) -> str:
    import docx2txt

    return docx2txt.process(file_path)


def _read_excel(file_path: str) -> str:
    """读取 Excel 文件（xlsx 需 openpyxl，xls 需 xlrd）。"""
    import pandas as pd

    excel_file = pd.ExcelFile(file_path)
    parts = []
    for sheet_name in excel_file.sheet_names:
        df = excel_file.parse(sheet_name)
        parts.append(f"=== {sheet_name} ===\n{df.to_string(index=False)}")
    return "\n\n".join(parts)


def _read_pptx(file_path: str) -> str:
    from pptx import Presentation

    prs = Presentation(file_path)
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
        if texts:
            slides.append("\n".join(texts))
    return "\n\n".join(slides)


_EXT_HANDLERS: Dict[str, Callable[[str], str]] = {
    "pdf": _read_pdf,
    "docx": _read_docx,
    "xlsx": _read_excel,
    "xls": _read_excel,
    "ods": _read_excel,
    "pptx": _read_pptx,
    "csv": _read_csv,
    "tsv": _read_csv,
}

# 无法解析但值得给出明确提示的旧格式
_UNSUPPORTED_HINTS = {
    "doc": "暂不支持旧版 .doc 格式，请在 Word 中另存为 .docx 后重新发送",
    "ppt": "暂不支持旧版 .ppt 格式，请在 PowerPoint 中另存为 .pptx 后重新发送",
    "zip": "暂不支持压缩包，请解压后发送其中的文件",
    "rar": "暂不支持压缩包，请解压后发送其中的文件",
    "7z": "暂不支持压缩包，请解压后发送其中的文件",
    "mp3": "不支持音频文件",
    "mp4": "不支持视频文件",
}


def supported_extensions() -> set[str]:
    """返回全部可解析的扩展名集合。"""
    return set(_EXT_HANDLERS) | TEXT_EXTENSIONS


def read_any_file_to_text(file_path: str, allowed: Optional[set[str]] = None) -> str:
    """将文件解析为文本。

    Args:
        file_path: 文件路径。
        allowed: 可接受的扩展名集合，None 表示不限制。

    Raises:
        FileParseError: 文件不存在、类型不支持或解析失败。
    """
    if isinstance(file_path, bytes):
        file_path = file_path.decode("utf-8", errors="replace")
    file_path = os.path.abspath(os.path.normpath(file_path))

    if not os.path.isfile(file_path):
        raise FileParseError(f"文件不存在: {file_path}")

    ext = detect_extension(file_path)

    if ext in _UNSUPPORTED_HINTS and (allowed is None or ext not in allowed):
        raise FileParseError(_UNSUPPORTED_HINTS[ext])
    if allowed is not None and ext not in allowed:
        raise FileParseError(f"不支持的文件类型: {ext or '(无扩展名)'}")

    try:
        if ext in _EXT_HANDLERS:
            text = _EXT_HANDLERS[ext](file_path)
        elif ext in TEXT_EXTENSIONS or allowed is None:
            text = _read_text(file_path)
        else:
            raise FileParseError(f"不支持的文件类型: {ext or '(无扩展名)'}")
    except FileParseError:
        raise
    except ImportError as exc:
        raise FileParseError(f"缺少解析依赖: {exc}，请检查插件 requirements 是否已安装") from exc
    except Exception as exc:
        raise FileParseError(f"读取 {ext or '文件'} 失败: {exc}") from exc

    if text is None:
        text = ""
    return text.strip()
